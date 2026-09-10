from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.shortcuts import get_object_or_404, redirect
from django.urls import reverse
from django.utils import timezone
from django.views.generic import DetailView, CreateView, UpdateView, ListView

from apps.purchase_orders.models import PurchaseOrder, ApprovalLog
from .models import ProcurementOrder, ProcurementLineItem
from .forms import ProcurementOrderForm


class ProcurementRequiredMixin(LoginRequiredMixin, UserPassesTestMixin):
    def test_func(self):
        return self.request.user.is_procurement_officer or self.request.user.is_staff


class ProcurementDashboardView(ProcurementRequiredMixin, ListView):
    template_name = 'procurement/dashboard.html'
    context_object_name = 'open_pos'

    def get_queryset(self):
        # ONLY fully head-approved POs (SENT_PROC) — nothing less
        return PurchaseOrder.objects.filter(
            status=PurchaseOrder.Status.SENT_PROC
        ).select_related('department', 'requester').order_by('-sent_to_procurement_at')

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['ordered'] = ProcurementOrder.objects.filter(
            status=ProcurementOrder.OrderStatus.ORDERED
        ).select_related('purchase_order', 'purchase_order__department').order_by('-ordered_at')
        return ctx


class ProcurementOrderCreateView(ProcurementRequiredMixin, CreateView):
    model = ProcurementOrder
    form_class = ProcurementOrderForm
    template_name = 'procurement/order_form.html'

    def dispatch(self, request, *args, **kwargs):
        self.po = get_object_or_404(
            PurchaseOrder, pk=self.kwargs['po_pk'],
            status=PurchaseOrder.Status.SENT_PROC)
        if hasattr(self.po, 'procurement_order'):
            return redirect(reverse('procurement:order_detail',
                                    kwargs={'pk': self.po.procurement_order.pk}))
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['po'] = self.po
        ctx['active_items'] = list(self.po.active_line_items)
        return ctx

    @transaction.atomic
    def form_valid(self, form):
        po = self.po
        proc = form.save(commit=False)
        proc.purchase_order = po
        proc.ordered_by = self.request.user
        proc.status = ProcurementOrder.OrderStatus.OPEN
        proc.save()

        for item in po.active_line_items:
            try:
                price = float(self.request.POST.get(f'price_{item.pk}', 0) or 0)
            except (ValueError, TypeError):
                price = 0
            ProcurementLineItem.objects.create(
                procurement_order=proc, po_line_item=item, unit_price=price)
            if price:
                item.unit_price = price
                item.save()

        messages.success(self.request, f'Procurement record opened for {po.po_number}.')
        return redirect(reverse('procurement:order_detail', kwargs={'pk': proc.pk}))


class ProcurementOrderDetailView(ProcurementRequiredMixin, DetailView):
    model = ProcurementOrder
    template_name = 'procurement/order_detail.html'
    context_object_name = 'proc_order'


class ProcurementOrderUpdateView(ProcurementRequiredMixin, UpdateView):
    model = ProcurementOrder
    form_class = ProcurementOrderForm
    template_name = 'procurement/order_form.html'

    def dispatch(self, request, *args, **kwargs):
        obj = get_object_or_404(ProcurementOrder, pk=self.kwargs['pk'])
        if obj.status == ProcurementOrder.OrderStatus.ORDERED:
            messages.warning(request, 'This order is already marked as ordered.')
            return redirect(reverse('procurement:order_detail', kwargs={'pk': obj.pk}))
        self.po = obj.purchase_order
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['po'] = self.po
        ctx['active_items'] = list(self.po.active_line_items)
        ctx['proc_line_map'] = {pl.po_line_item_id: pl for pl in self.object.proc_line_items.all()}
        ctx['is_edit'] = True
        return ctx

    @transaction.atomic
    def form_valid(self, form):
        proc = form.save()
        for item in self.po.active_line_items:
            try:
                price = float(self.request.POST.get(f'price_{item.pk}', 0) or 0)
            except (ValueError, TypeError):
                price = 0
            pl = proc.proc_line_items.filter(po_line_item=item).first()
            if pl:
                pl.unit_price = price
                pl.save()
            if price:
                item.unit_price = price
                item.save()
        messages.success(self.request, 'Procurement record updated.')
        return redirect(reverse('procurement:order_detail', kwargs={'pk': proc.pk}))


@login_required
@transaction.atomic
def mark_ordered(request, pk):
    """Mark as ORDERED → immediately available to stores for receiving."""
    if not (request.user.is_procurement_officer or request.user.is_staff):
        raise PermissionDenied
    proc = get_object_or_404(ProcurementOrder, pk=pk)
    if request.method == 'POST':
        if proc.status != ProcurementOrder.OrderStatus.OPEN:
            messages.error(request, 'This order has already been marked as ordered.')
        else:
            proc.status = ProcurementOrder.OrderStatus.ORDERED
            proc.ordered_at = timezone.now()
            proc.save()

            po = proc.purchase_order
            prev = po.status
            po.status = PurchaseOrder.Status.ORDERED
            po.save()

            ApprovalLog.objects.create(
                purchase_order=po, actor=request.user,
                from_status=prev, to_status=po.status,
                comment='Marked as ordered by procurement — sent to stores for receiving.',
            )
            from apps.notifications.utils import notify_order_placed, notify_stores_on_order
            notify_order_placed(po, proc)
            notify_stores_on_order(po)
            messages.success(request, f'{po.po_number} marked as ordered and sent to stores.')
    return redirect(reverse('procurement:order_detail', kwargs={'pk': pk}))


@login_required
def export_orders_excel(request):
    """Export procurement orders for a date range as Excel."""
    if not (request.user.is_procurement_officer or request.user.is_staff):
        raise PermissionDenied
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from io import BytesIO
    from django.utils import timezone as tz
    date_from=request.GET.get("from",""); date_to=request.GET.get("to","")
    qs=ProcurementOrder.objects.select_related(
        "purchase_order","purchase_order__department","purchase_order__requester","ordered_by"
    ).order_by("-ordered_at")
    if date_from:
        try:
            from datetime import datetime
            qs=qs.filter(ordered_at__gte=datetime.strptime(date_from,"%Y-%m-%d"))
        except ValueError: pass
    if date_to:
        try:
            from datetime import datetime
            qs=qs.filter(ordered_at__lte=datetime.strptime(date_to+" 23:59:59","%Y-%m-%d %H:%M:%S"))
        except ValueError: pass
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Procurement Orders"
    BRAND="FF1B3F6E"
    heads=["PO Number","Title","Department","Requester","Supplier","Order Ref",
           "Ordered By","Ordered At","Expected Delivery","Status","Total Amount"]
    widths=[16,30,16,22,24,18,22,18,18,14,16]
    for ci,(h,ww) in enumerate(zip(heads,widths),1):
        c=ws.cell(row=1,column=ci,value=h)
        c.font=Font(bold=True,color="FFFFFFFF",name="Arial",size=9)
        c.fill=PatternFill("solid",start_color=BRAND,fgColor=BRAND)
        c.alignment=Alignment(horizontal="center",vertical="center")
        ws.column_dimensions[get_column_letter(ci)].width=ww
    ws.row_dimensions[1].height=22
    for ri,o in enumerate(qs,2):
        po=o.purchase_order
        row=[po.po_number,po.title,
             po.department.name if po.department else "",
             po.requester.get_full_name() if po.requester else "",
             o.supplier_name or "",o.order_reference or "",
             o.ordered_by.get_full_name() if o.ordered_by else "",
             o.ordered_at.strftime("%d/%m/%Y %H:%M") if o.ordered_at else "",
             o.expected_delivery.strftime("%d/%m/%Y") if o.expected_delivery else "",
             o.get_status_display(),float(po.total_amount)]
        for ci,val in enumerate(row,1): ws.cell(row=ri,column=ci,value=val)
        if ri%2==0:
            for ci in range(1,len(heads)+1):
                ws.cell(row=ri,column=ci).fill=PatternFill("solid",start_color="FFF8FAFC",fgColor="FFF8FAFC")
    buf=BytesIO(); wb.save(buf); buf.seek(0)
    fn=f'Procurement_Orders_{tz.now().strftime("%Y%m%d")}.xlsx'
    resp=HttpResponse(buf.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"]=f'attachment; filename="{fn}"'
    return resp
