from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db.models import Q, Count, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect
from django.urls import reverse
from django.utils import timezone
from django.views.generic import ListView, DetailView, CreateView, UpdateView, TemplateView
from .models import Asset, AssetCategory, AssetMutation
from .forms import AssetForm, AllocateForm, TransferForm, OffboardForm, StatusChangeForm, MutationNoteForm

def _can_manage(user):
    return user.is_staff or user.is_head_approver or getattr(user,"is_asset_manager",False)

def _can_view(user, asset):
    if _can_manage(user): return True
    if asset.allocated_to==user: return True
    if user.is_hod and user.department and asset.department==user.department: return True
    return False

class AMixin(LoginRequiredMixin, UserPassesTestMixin):
    def test_func(self): return _can_manage(self.request.user)

class AssetDashboardView(LoginRequiredMixin, TemplateView):
    template_name="assets/dashboard.html"
    def get_context_data(self,**kw):
        ctx=super().get_context_data(**kw); u=self.request.user
        if _can_manage(u): base=Asset.objects.all()
        elif u.is_hod and u.department: base=Asset.objects.filter(department=u.department)
        else: base=Asset.objects.filter(allocated_to=u)
        ctx["total"]=base.count()
        ctx["active"]=base.filter(status="ACTIVE").count()
        ctx["under_repair"]=base.filter(status="UNDER_REPAIR").count()
        ctx["unallocated"]=base.filter(status="ACTIVE",allocated_to__isnull=True).count()
        ctx["total_cost"]=base.filter(purchase_cost__isnull=False).aggregate(s=Sum("purchase_cost"))["s"] or 0
        ctx["by_category"]=list(base.values("category__name").annotate(count=Count("pk")).order_by("-count"))
        if _can_manage(u):
            ctx["by_department"]=list(base.filter(department__isnull=False).values("department__name").annotate(count=Count("pk")).order_by("-count"))
        ctx["my_assets"]=Asset.objects.filter(allocated_to=u,status="ACTIVE").select_related("category","department")
        ctx["recent_mutations"]=AssetMutation.objects.filter(asset__in=base).select_related("asset","actor","to_user","from_user")[:10]
        in30=timezone.now().date()+timezone.timedelta(days=30)
        ctx["warranty_expiring"]=base.filter(warranty_expiry__isnull=False,warranty_expiry__lte=in30,warranty_expiry__gte=timezone.now().date(),status="ACTIVE").order_by("warranty_expiry")[:8]
        ctx["can_manage"]=_can_manage(u)
        return ctx

class AssetListView(LoginRequiredMixin, ListView):
    model=Asset; template_name="assets/list.html"; context_object_name="assets"; paginate_by=30
    def get_queryset(self):
        u=self.request.user
        if _can_manage(u): qs=Asset.objects.all()
        elif u.is_hod and u.department: qs=Asset.objects.filter(department=u.department)
        else: qs=Asset.objects.filter(allocated_to=u)
        qs=qs.select_related("category","department","allocated_to")
        q=self.request.GET.get("q",""); stat=self.request.GET.get("status","")
        cat=self.request.GET.get("category",""); dept=self.request.GET.get("department","")
        alloc=self.request.GET.get("allocated","")
        if q: qs=qs.filter(Q(asset_tag__icontains=q)|Q(name__icontains=q)|Q(serial_number__icontains=q))
        if stat: qs=qs.filter(status=stat)
        if cat: qs=qs.filter(category_id=cat)
        if dept: qs=qs.filter(department_id=dept)
        if alloc=="yes": qs=qs.filter(allocated_to__isnull=False)
        elif alloc=="no": qs=qs.filter(allocated_to__isnull=True)
        return qs
    def get_context_data(self,**kw):
        ctx=super().get_context_data(**kw)
        ctx.update({"q":self.request.GET.get("q",""),"status_f":self.request.GET.get("status",""),
            "cat_f":self.request.GET.get("category",""),"dept_f":self.request.GET.get("department",""),
            "alloc_f":self.request.GET.get("allocated",""),"categories":AssetCategory.objects.all(),
            "status_choices":Asset.Status.choices,"can_manage":_can_manage(self.request.user)})
        from apps.departments.models import Department
        ctx["departments"]=Department.objects.all()
        return ctx

class AssetDetailView(LoginRequiredMixin, DetailView):
    model=Asset; template_name="assets/detail.html"; context_object_name="asset"
    def get_object(self):
        a=get_object_or_404(Asset,pk=self.kwargs["pk"])
        if not _can_view(self.request.user,a): raise PermissionDenied
        return a
    def get_context_data(self,**kw):
        ctx=super().get_context_data(**kw); u=self.request.user; a=self.object
        ctx["can_manage"]=_can_manage(u)
        ctx["can_offboard"]=a.allocated_to and (_can_manage(u) or (u.is_hod and u.department==a.department))
        ctx["can_transfer"]=_can_manage(u) or (u.is_hod and u.department==a.department)
        ctx["mutations"]=a.mutations.select_related("actor","from_user","to_user","from_department","to_department")
        ctx["allocate_form"]=AllocateForm(); ctx["transfer_form"]=TransferForm()
        ctx["offboard_form"]=OffboardForm()
        ctx["status_form"]=StatusChangeForm(initial={"status":a.status,"condition":a.condition})
        ctx["note_form"]=MutationNoteForm()
        return ctx

class AssetCreateView(AMixin, CreateView):
    model=Asset; form_class=AssetForm; template_name="assets/form.html"
    def form_valid(self,form):
        a=form.save(commit=False); a.created_by=self.request.user
        if a.allocated_to: a.allocated_at=timezone.now()
        a.save()
        if a.allocated_to:
            AssetMutation.objects.create(asset=a,mutation_type="ALLOCATED",actor=self.request.user,
                to_user=a.allocated_to,to_department=a.department,notes="Initial allocation.")
        messages.success(self.request,f"Asset {a.asset_tag} created.")
        return redirect(reverse("assets:detail",kwargs={"pk":a.pk}))
    def get_context_data(self,**kw):
        ctx=super().get_context_data(**kw); ctx["page_title"]="Add Asset"; return ctx

class AssetUpdateView(AMixin, UpdateView):
    model=Asset; form_class=AssetForm; template_name="assets/form.html"
    def form_valid(self,form):
        a=form.save(); messages.success(self.request,f"Asset {a.asset_tag} updated.")
        return redirect(reverse("assets:detail",kwargs={"pk":a.pk}))
    def get_context_data(self,**kw):
        ctx=super().get_context_data(**kw); ctx["page_title"]=f"Edit {self.object.asset_tag}"; return ctx

@login_required
def allocate_asset(request,pk):
    a=get_object_or_404(Asset,pk=pk)
    if not(_can_manage(request.user) or (request.user.is_hod and request.user.department==a.department)): raise PermissionDenied
    if request.method=="POST":
        form=AllocateForm(request.POST)
        if form.is_valid():
            prev=a.allocated_to; a.allocated_to=form.cleaned_data["allocated_to"]; a.allocated_at=timezone.now()
            a.save(update_fields=["allocated_to","allocated_at"])
            AssetMutation.objects.create(asset=a,mutation_type="ALLOCATED",actor=request.user,
                from_user=prev,to_user=a.allocated_to,to_department=a.department,notes=form.cleaned_data.get("notes",""))
            messages.success(request,f"{a.asset_tag} allocated to {a.allocated_to.get_full_name()}.")
    return redirect(reverse("assets:detail",kwargs={"pk":pk}))

@login_required
def transfer_asset(request,pk):
    a=get_object_or_404(Asset,pk=pk)
    if not(_can_manage(request.user) or (request.user.is_hod and request.user.department==a.department)): raise PermissionDenied
    if request.method=="POST":
        form=TransferForm(request.POST)
        if form.is_valid():
            prev=a.department; new=form.cleaned_data["to_department"]
            a.department=new; a.save(update_fields=["department"])
            AssetMutation.objects.create(asset=a,mutation_type="TRANSFERRED",actor=request.user,
                from_department=prev,to_department=new,notes=form.cleaned_data.get("notes",""))
            messages.success(request,f"{a.asset_tag} transferred to {new.name}.")
    return redirect(reverse("assets:detail",kwargs={"pk":pk}))

@login_required
def offboard_asset(request,pk):
    a=get_object_or_404(Asset,pk=pk)
    if not(_can_manage(request.user) or (request.user.is_hod and request.user.department==a.department)): raise PermissionDenied
    if request.method=="POST":
        form=OffboardForm(request.POST)
        if form.is_valid():
            pu=a.allocated_to; ps=a.status; ns=form.cleaned_data["new_status"]
            a.allocated_to=None; a.allocated_at=None; a.status=ns
            a.save(update_fields=["allocated_to","allocated_at","status"])
            AssetMutation.objects.create(asset=a,mutation_type="OFFBOARDED",actor=request.user,
                from_user=pu,from_status=ps,to_status=ns,notes=form.cleaned_data["notes"])
            messages.success(request,f"{a.asset_tag} offboarded.")
    return redirect(reverse("assets:detail",kwargs={"pk":pk}))

@login_required
def change_status(request,pk):
    a=get_object_or_404(Asset,pk=pk)
    if not(_can_manage(request.user) or (request.user.is_hod and request.user.department==a.department)): raise PermissionDenied
    if request.method=="POST":
        form=StatusChangeForm(request.POST)
        if form.is_valid():
            prev=a.status; a.status=form.cleaned_data["status"]; a.condition=form.cleaned_data["condition"]
            a.save(update_fields=["status","condition"])
            AssetMutation.objects.create(asset=a,mutation_type="STATUS_CHANGE",actor=request.user,
                from_status=prev,to_status=a.status,notes=form.cleaned_data.get("notes",""))
            messages.success(request,f"{a.asset_tag} status updated.")
    return redirect(reverse("assets:detail",kwargs={"pk":pk}))

@login_required
def add_note(request,pk):
    a=get_object_or_404(Asset,pk=pk)
    if not _can_view(request.user,a): raise PermissionDenied
    if request.method=="POST":
        form=MutationNoteForm(request.POST)
        if form.is_valid():
            AssetMutation.objects.create(asset=a,mutation_type="NOTE",actor=request.user,notes=form.cleaned_data["notes"])
            messages.success(request,"Note added.")
    return redirect(reverse("assets:detail",kwargs={"pk":pk}))

@login_required
def export_assets(request):
    if not _can_manage(request.user): raise PermissionDenied
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    from io import BytesIO
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Asset Register"
    BRAND="FF1B3F6E"
    heads=["Asset Tag","Name","Category","Brand","Model","Serial No.","Status","Condition",
           "Department","Allocated To","Location","Purchase Date","Purchase Cost",
           "Depreciated Value","Warranty Expiry","Notes"]
    widths=[14,28,16,14,16,18,12,12,16,24,16,14,14,16,14,30]
    for ci,(h,ww) in enumerate(zip(heads,widths),1):
        c=ws.cell(row=1,column=ci,value=h)
        c.font=Font(bold=True,color="FFFFFFFF",name="Arial",size=9)
        c.fill=PatternFill("solid",start_color=BRAND,fgColor=BRAND)
        c.alignment=Alignment(horizontal="center",vertical="center")
        ws.column_dimensions[get_column_letter(ci)].width=ww
    ws.row_dimensions[1].height=22
    for ri,a in enumerate(Asset.objects.select_related("category","department","allocated_to").order_by("asset_tag"),2):
        row=[a.asset_tag,a.name,a.category.name if a.category else "",a.brand,a.model_number,
             a.serial_number,a.get_status_display(),a.get_condition_display(),
             a.department.name if a.department else "",
             a.allocated_to.get_full_name() if a.allocated_to else "",a.location,
             a.purchase_date.strftime("%d/%m/%Y") if a.purchase_date else "",
             float(a.purchase_cost) if a.purchase_cost else "",
             a.depreciated_value if a.depreciated_value is not None else "",
             a.warranty_expiry.strftime("%d/%m/%Y") if a.warranty_expiry else "",a.notes]
        for ci,val in enumerate(row,1): ws.cell(row=ri,column=ci,value=val)
        if ri%2==0:
            for ci in range(1,len(heads)+1):
                ws.cell(row=ri,column=ci).fill=PatternFill("solid",start_color="FFF8FAFC",fgColor="FFF8FAFC")
    buf=BytesIO(); wb.save(buf); buf.seek(0)
    resp=HttpResponse(buf.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"]=f'attachment; filename="Asset_Register_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    return resp
