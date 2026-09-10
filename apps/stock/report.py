import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.utils import timezone

C_BRAND='1B3F6E'; C_WHITE='FFFFFF'; C_BORDER='BFBFBF'
THIN=Side(style='thin',color=C_BORDER)
TB=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)
MONTHS=['JUN','JUL','AUG','SEP','OCT','NOV','DEC','JAN','FEB','MAR','APR','MAY']
MFIELDS=['usage_jun','usage_jul','usage_aug','usage_sep','usage_oct','usage_nov',
         'usage_dec','usage_jan','usage_feb','usage_mar','usage_apr','usage_may']
HEADERS=['QB CODE','ITEM DESCRIPTION','U/M','CATEGORY',*MONTHS,
         'MONTHLY AVG','ANNUAL TOTAL','CURRENT STOCK','STOCK COVER (MO)',
         'OPEN ORDER','REORDER QTY','MIN THRESHOLD','MAX (ANNUAL)','ACTION',
         'LEAD TIME','TO ORDER?','ORDER QTY']

def _fill(h): return PatternFill('solid',start_color=h,fgColor=h)
def _action_fill(a):
    return {'OUT OF STOCK':_fill('FFCCCC'),'REORDER NOW':_fill('FFCCCC'),
            'REORDER SOON':_fill('FFF2CC'),'OVERSTOCKED':_fill('DDEEFF'),
            'WELL STOCKED':_fill('E2EFDA')}.get(a,_fill('F2F2F2'))

def generate_stock_report(items):
    wb=Workbook(); ws=wb.active; ws.title='Stock Analysis Report'
    n=len(HEADERS)
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n)
    c=ws.cell(row=1,column=1,value=f'STOCK ANALYSIS REPORT — {timezone.now().strftime("%d %b %Y %H:%M")}')
    c.font=Font(bold=True,size=13,color=C_WHITE,name='Arial')
    c.fill=_fill(C_BRAND); c.alignment=Alignment(horizontal='center',vertical='center')
    ws.row_dimensions[1].height=24
    for ci,h in enumerate(HEADERS,1):
        cell=ws.cell(row=2,column=ci,value=h)
        cell.font=Font(bold=True,size=9,color=C_WHITE,name='Arial')
        cell.fill=_fill(C_BRAND); cell.border=TB
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.row_dimensions[2].height=40
    ws.column_dimensions['A'].width=14; ws.column_dimensions['B'].width=30
    ws.column_dimensions['C'].width=8; ws.column_dimensions['D'].width=16
    for ci in range(5,n+1): ws.column_dimensions[get_column_letter(ci)].width=11
    for ri,item in enumerate(items,3):
        a=item.action; rf=_action_fill(a)
        vals=[item.qb_code,item.description,item.uom,item.category or '',
              *[float(getattr(item,f)) for f in MFIELDS],
              item.monthly_moving_average,item.annual_usage,
              float(item.current_stock),
              item.stock_cover_months if item.stock_cover_months is not None else '',
              float(item.open_order),item.target_reorder_value,
              item.min_threshold,item.max_threshold,a,
              float(item.lead_time_months),
              'YES' if item.to_order else 'NO',float(item.order_qty)]
        for ci,val in enumerate(vals,1):
            cell=ws.cell(row=ri,column=ci,value=val)
            cell.border=TB; cell.fill=rf
            cell.font=Font(size=9,name='Arial',bold=(HEADERS[ci-1]=='ACTION'))
            cell.alignment=Alignment(vertical='center',horizontal='left' if ci<=4 else 'center')
            if ci>4 and isinstance(val,(int,float)):
                cell.number_format='#,##0.00'
    sc=n+2
    ws.cell(row=1,column=sc,value='SUMMARY').font=Font(bold=True,name='Arial')
    for ri2,(label,count) in enumerate([
        ('Total Items',len(items)),
        ('Out of Stock',sum(1 for i in items if i.action=='OUT OF STOCK')),
        ('Reorder Now',sum(1 for i in items if i.action=='REORDER NOW')),
        ('Reorder Soon',sum(1 for i in items if i.action=='REORDER SOON')),
        ('Overstocked',sum(1 for i in items if i.action=='OVERSTOCKED')),
        ('Well Stocked',sum(1 for i in items if i.action=='WELL STOCKED')),
        ('No Usage Data',sum(1 for i in items if i.action=='NO USAGE DATA')),
    ],2):
        lc=ws.cell(row=ri2,column=sc,value=label)
        vc=ws.cell(row=ri2,column=sc+1,value=count)
        lc.font=vc.font=Font(size=9,name='Arial')
        lc.border=vc.border=TB
        vc.alignment=Alignment(horizontal='center')
    ws.column_dimensions[get_column_letter(sc)].width=18
    ws.column_dimensions[get_column_letter(sc+1)].width=10
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf
