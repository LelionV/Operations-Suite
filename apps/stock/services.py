import openpyxl, csv, io
from decimal import Decimal, InvalidOperation
from .models import StockItem, StockUploadLog

def _dec(val, default=0):
    try:
        if val is None or str(val).strip() in ('','-','N/A'): return Decimal(default)
        return Decimal(str(val).strip())
    except InvalidOperation:
        return Decimal(default)

def _parse_xlsx(ws):
    header = [str(c.value).strip().upper() if c.value else '' for c in ws[2]]
    def col(row, name):
        try: return row[header.index(name)]
        except (ValueError, IndexError): return None
    for row in ws.iter_rows(min_row=3, values_only=True):
        code = col(row,'QB CODE')
        if not code or not str(code).strip(): continue
        yield {
            'qb_code': str(code).strip(),
            'description': str(col(row,'ITEM DESCRIPTION') or '').strip(),
            'uom': str(col(row,'U/M') or '').strip(),
            'current_stock': _dec(col(row,'STOCK STATUS AS AT XX-XX-2026')),
            'open_order': _dec(col(row,'OPEN ORDER')),
            'lead_time_months': _dec(col(row,'LEADTIME\n(MONTHS)'),3),
            'category': str(col(row,'CATEGORY') or '').strip(),
            'to_order': str(col(row,'FEEDBACK - TO ORDER YES/NO') or '').upper()=='YES',
            'order_qty': _dec(col(row,'FEEDBACK - QUANTITY')),
            'usage_jun':_dec(col(row,'JUN')),'usage_jul':_dec(col(row,'JUL')),
            'usage_aug':_dec(col(row,'AUG')),'usage_sep':_dec(col(row,'SEP')),
            'usage_oct':_dec(col(row,'OCT')),'usage_nov':_dec(col(row,'NOV')),
            'usage_dec':_dec(col(row,'DEC')),'usage_jan':_dec(col(row,'JAN')),
            'usage_feb':_dec(col(row,'FEB')),'usage_mar':_dec(col(row,'MAR')),
            'usage_apr':_dec(col(row,'APR')),'usage_may':_dec(col(row,'MAY')),
        }

def process_upload(file_obj, filename, user):
    rows = []
    ext = filename.lower().split('.')[-1]
    if ext in ('xlsx','xlsm'):
        wb = openpyxl.load_workbook(file_obj, data_only=True)
        rows = list(_parse_xlsx(wb.active))
    created=updated=skipped=0
    for data in rows:
        code = data.pop('qb_code')
        if not code: skipped+=1; continue
        if not data.get('lead_time_months') or data['lead_time_months']<=0:
            data['lead_time_months']=Decimal('3')
        _, was_created = StockItem.objects.update_or_create(qb_code=code, defaults=data)
        if was_created: created+=1
        else: updated+=1
    log = StockUploadLog.objects.create(
        filename=filename, uploaded_by=user,
        rows_total=len(rows), rows_created=created,
        rows_updated=updated, rows_skipped=skipped)
    return log, created, updated, skipped
